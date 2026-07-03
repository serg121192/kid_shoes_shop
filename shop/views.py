import io
import json
import logging
from datetime import timedelta

from django.conf import settings
from django.contrib.auth import get_user_model
from django.db import transaction
from django.db.models import (
    Count,
    Sum,
    F,
    ExpressionWrapper,
    DecimalField as DjDecimalField,
)
from django.core.exceptions import ValidationError
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone

from django_filters.rest_framework import DjangoFilterBackend
from rest_framework.decorators import action, api_view, permission_classes
from rest_framework.filters import SearchFilter, OrderingFilter
from rest_framework.request import Request
from rest_framework.response import Response
from rest_framework import viewsets, status
from rest_framework.permissions import IsAuthenticated, AllowAny, IsAdminUser

from shop.emails import (
    send_order_confirmation,
    send_order_status_update,
    send_new_order_alert,
)
from shop.pagination import CatalogPagination
from shop.r2_upload import (
    is_r2_direct_upload_enabled,
    presign_gallery_uploads,
    presign_video_upload,
    validate_gallery_key,
    validate_video_key,
    object_exists,
)
from shop.permissions import (
    IsAdminOrReadOnly,
    IsAdminOrVideoUploadToken,
    IsOwnerOrAdmin,
    IsStaffOrSeller,
    can_manage_orders,
)
from shop.video_upload import make_video_upload_token, verify_video_upload_token
from shop.filters import ProductFilter
from shop.models import (
    Product,
    ProductSize,
    ProductImage,
    ProductVideo,
    Vendor,
    Cart,
    CartItem,
    Order,
    OrderItem,
    Wishlist,
    WishlistItem,
    DeliveryInfo,
    Review,
    SiteVisit,
)
from shop.nova_poshta import (
    get_cities_list,
    get_warehouses_list,
    get_tracking_status,
)
from shop.serializers import (
    ProductSerializer,
    ProductListSerializer,
    ProductRetrieveSerializer,
    ProductImageSerializer,
    ProductVideoSerializer,
    ProductSizeListSerializer,
    VendorSerializer,
    CartSerializer,
    AddToCartSerializer,
    RemoveFromCartSerializer,
    WishlistSerializer,
    WishlistItemSerializer,
    AddToWishlistSerializer,
    RemoveFromWishlistSerializer,
    OrderSerializer,
    OrderStatusSerializer,
    StaffCreateOrderSerializer,
    QuickSaleSerializer,
    OrderItemSerializer,
    DeliveryInfoSerializer,
    ReviewSerializer,
)

logger = logging.getLogger(__name__)

User = get_user_model()


def _create_order_with_items(
    *,
    order_user,
    delivery_data: dict,
    items_data: list[dict],
    created_by=None,
    initial_status=None,
    send_emails: bool = True,
    sale_channel=Order.SaleChannelChoices.ONLINE,
) -> Order:
    """Create order, delivery, items; reduce stock. Caller must wrap in transaction."""
    order = Order.objects.create(
        user=order_user,
        created_by=created_by,
        status=initial_status or Order.StatusChoices.PENDING,
        sale_channel=sale_channel,
    )
    DeliveryInfo.objects.create(order=order, **delivery_data)

    total_price = 0
    for item in items_data:
        product_size = ProductSize.objects.select_related("product").select_for_update().get(
            id=item["product_size"]
        )
        quantity = item["quantity"]
        try:
            OrderItem.validate_product_quantity(
                product_size=product_size,
                quantity=quantity,
                error_to_raise=ValidationError,
            )
        except ValidationError as exc:
            raise ValidationError(exc.message) from exc

        OrderItem.objects.create(
            order=order,
            product_size=product_size,
            quantity=quantity,
            price=product_size.product.discounted_price,
        )
        total_price += quantity * product_size.product.discounted_price
        product_size.reduce_stock(quantity)

    order.total_price = total_price
    order.save()

    if send_emails:
        def _send_emails():
            try:
                if order.user_id:
                    send_order_confirmation(order)
                send_new_order_alert(order)
            except Exception:
                logger.exception(
                    "Order #%s: failed to send notification emails", order.pk
                )

        transaction.on_commit(_send_emails)

    return order


class ProductViewSet(viewsets.ModelViewSet):
    queryset = Product.objects.all()
    permission_classes = [IsAdminOrReadOnly]
    pagination_class = CatalogPagination
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_class = ProductFilter
    search_fields = ["model_name", "vendor__name", "description"]
    ordering_fields = ["full_price", "discount"]
    ordering = ["full_price"]
    lookup_field = "slug"
    lookup_value_regex = "[^/]+"

    def get_object(self):
        lookup_url_kwarg = self.lookup_url_kwarg or self.lookup_field
        lookup_value = self.kwargs.get(lookup_url_kwarg)
        queryset = self.filter_queryset(self.get_queryset())

        try:
            obj = get_object_or_404(
                queryset, **{self.lookup_field: lookup_value}
            )
        except Http404:
            if lookup_value and str(lookup_value).isdigit():
                obj = get_object_or_404(queryset, pk=lookup_value)
            else:
                raise

        self.check_object_permissions(self.request, obj)
        return obj

    def get_queryset(self):
        qs = Product.objects.select_related("vendor")
        if self.action == "list":
            qs = qs.prefetch_related("sizes", "images").annotate(
                _total_qty=Sum("sizes__quantity")
            )
        else:
            qs = qs.prefetch_related("sizes", "images", "videos")
        if self.request.query_params.get("size"):
            qs = qs.distinct()
        user = self.request.user
        if self.action in ("list", "retrieve") and not (
            user.is_authenticated and user.is_staff
        ):
            qs = qs.filter(is_published=True)
        return qs

    def get_serializer_class(self):
        if self.action == "list":
            return ProductListSerializer
        if self.action == "retrieve":
            return ProductRetrieveSerializer
        return ProductSerializer

    def list(self, request, *args, **kwargs):
        queryset = self.filter_queryset(self.get_queryset())
        page = self.paginate_queryset(queryset)
        serializer = self.get_serializer(page, many=True)
        response = self.get_paginated_response(serializer.data)
        user = request.user
        if user.is_authenticated and user.is_staff:
            response.data["ready_to_publish_count"] = Product.objects.filter(
                is_published=False, full_price__gt=0
            ).count()
        return response

    @action(
        detail=False,
        methods=["post"],
        url_path="bulk_publish",
        permission_classes=[IsAdminUser],
    )
    def bulk_publish(self, request: Request) -> Response:
        """Publish products that have a price. Optional body: {"ids": [1, 2, 3]}."""
        ids = request.data.get("ids")
        qs = Product.objects.filter(full_price__gt=0, is_published=False)
        if ids is not None:
            if not isinstance(ids, list) or not ids:
                return Response(
                    {"error": "ids must be a non-empty list"},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            qs = qs.filter(id__in=ids)
        published = qs.update(is_published=True)
        return Response({"published": published})

    @action(
        detail=False,
        methods=["get"],
        url_path="check_duplicate",
        permission_classes=[IsAdminUser],
    )
    def check_duplicate(self, request: Request) -> Response:
        vendor_id = request.query_params.get("vendor")
        model_name = (request.query_params.get("model_name") or "").strip()
        exclude_id = request.query_params.get("exclude")

        if not vendor_id or not model_name:
            return Response({"exists": False})

        qs = Product.objects.filter(
            vendor_id=vendor_id, model_name=model_name
        )
        if exclude_id:
            qs = qs.exclude(pk=exclude_id)

        product = qs.first()
        return Response(
            {
                "exists": product is not None,
                "product_id": product.id if product else None,
            }
        )

    @action(
        detail=True,
        methods=["post"],
        url_path="upload_image",
        permission_classes=[IsAdminUser],
    )
    def upload_image(
        self, request: Request, pk=None, slug=None, **kwargs
    ) -> Response:
        product = self.get_object()
        image = request.FILES.get("image")
        if not image:
            return Response(
                {"error": "Файл не завантажено"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        is_main = request.data.get("is_main", "false").lower() == "true"
        order_num = int(request.data.get("order", 0))
        pi = ProductImage.objects.create(
            product=product, image=image, is_main=is_main, order=order_num
        )
        return Response(
            ProductImageSerializer(pi, context={"request": request}).data,
            status=status.HTTP_201_CREATED,
        )

    @action(
        detail=True,
        methods=["post"],
        url_path="presign_images",
        permission_classes=[IsAdminUser],
    )
    def presign_images(
        self, request: Request, pk=None, slug=None, **kwargs
    ) -> Response:
        if not is_r2_direct_upload_enabled():
            return Response({"direct_upload": False, "reason": "r2_not_configured"})

        product = self.get_object()
        items = request.data.get("images")
        if not isinstance(items, list) or not items:
            return Response(
                {"error": "Передайте масив images з content_type"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            uploads = presign_gallery_uploads(product, items)
        except Exception:
            logging.exception("presign_images failed")
            return Response({"direct_upload": False})

        return Response({"direct_upload": True, "uploads": uploads})

    @action(
        detail=True,
        methods=["post"],
        url_path="confirm_images",
        permission_classes=[IsAdminUser],
    )
    def confirm_images(
        self, request: Request, pk=None, slug=None, **kwargs
    ) -> Response:
        if not is_r2_direct_upload_enabled():
            return Response(
                {"error": "Пряме завантаження недоступне"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        product = self.get_object()
        items = request.data.get("images")
        if not isinstance(items, list) or not items:
            return Response(
                {"error": "Передайте масив images з key, is_main, order"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        order_base = product.images.count()
        has_main = product.images.filter(is_main=True).exists()
        main_assigned = has_main
        created: list[ProductImage] = []

        for idx, item in enumerate(items):
            key = item.get("key")
            if not key or not validate_gallery_key(key, product):
                return Response(
                    {"error": f"Невірний ключ файлу: {key}"},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            if not object_exists(key):
                return Response(
                    {"error": f"Файл не знайдено в сховищі: {key}"},
                    status=status.HTTP_400_BAD_REQUEST,
                )

            is_main = bool(item.get("is_main", False))
            order_num = int(item.get("order", order_base + idx))
            set_as_main = not main_assigned and (is_main or idx == 0)
            if set_as_main:
                main_assigned = True
                ProductImage.objects.filter(
                    product=product, is_main=True
                ).update(is_main=False)

            pi = ProductImage(product=product, is_main=set_as_main, order=order_num)
            pi.image.name = key
            pi.image._committed = True
            pi.save()
            created.append(pi)

        return Response(
            ProductImageSerializer(
                created, many=True, context={"request": request}
            ).data,
            status=status.HTTP_201_CREATED,
        )

    @action(
        detail=True,
        methods=["post"],
        url_path="upload_images",
        permission_classes=[IsAdminUser],
    )
    def upload_images(
        self, request: Request, pk=None, slug=None, **kwargs
    ) -> Response:
        product = self.get_object()
        files = request.FILES.getlist("images")
        if not files:
            return Response(
                {"error": "Файли не завантажено"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        meta_raw = request.data.get("meta", "[]")
        try:
            meta = json.loads(meta_raw) if isinstance(meta_raw, str) else meta_raw
        except (json.JSONDecodeError, TypeError):
            meta = []

        order_base = product.images.count()
        has_main = product.images.filter(is_main=True).exists()
        main_assigned = has_main
        created: list[ProductImage] = []

        for idx, image in enumerate(files):
            item_meta = meta[idx] if idx < len(meta) else {}
            is_main = bool(item_meta.get("is_main", False))
            order_num = int(item_meta.get("order", order_base + idx))
            set_as_main = not main_assigned and (is_main or idx == 0)
            if set_as_main:
                main_assigned = True
                ProductImage.objects.filter(
                    product=product, is_main=True
                ).update(is_main=False)
            pi = ProductImage.objects.create(
                product=product,
                image=image,
                is_main=set_as_main,
                order=order_num,
            )
            created.append(pi)

        return Response(
            ProductImageSerializer(
                created, many=True, context={"request": request}
            ).data,
            status=status.HTTP_201_CREATED,
        )

    @action(
        detail=True,
        methods=["delete"],
        url_path=r"delete_image/(?P<image_id>[^/.]+)",
        permission_classes=[IsAdminUser],
    )
    def delete_image(
        self, request: Request, pk=None, image_id=None, slug=None, **kwargs
    ) -> Response:
        product = self.get_object()
        try:
            ProductImage.objects.get(pk=image_id, product=product).delete()
            return Response(status=status.HTTP_204_NO_CONTENT)
        except ProductImage.DoesNotExist:
            return Response(status=status.HTTP_404_NOT_FOUND)

    @action(
        detail=True,
        methods=["patch"],
        url_path=r"set_main_image/(?P<image_id>[^/.]+)",
        permission_classes=[IsAdminUser],
    )
    def set_main_image(
        self, request: Request, pk=None, image_id=None, slug=None, **kwargs
    ) -> Response:
        product = self.get_object()
        try:
            img = ProductImage.objects.get(pk=image_id, product=product)
            img.is_main = True
            img.save()  # model.save() автоматично скидає is_main у решті
            return Response(
                ProductImageSerializer(img, context={"request": request}).data
            )
        except ProductImage.DoesNotExist:
            return Response(status=status.HTTP_404_NOT_FOUND)

    @action(
        detail=True,
        methods=["post"],
        url_path="prepare_video_upload",
        permission_classes=[IsAdminUser],
    )
    def prepare_video_upload(
        self, request: Request, pk=None, slug=None, **kwargs
    ) -> Response:
        product = self.get_object()
        content_type = request.data.get("content_type") or "video/mp4"
        filename = request.data.get("filename") or ""

        backend_base = getattr(settings, "BACKEND_PUBLIC_URL", "").rstrip("/")
        backend_url = (
            f"{backend_base}/api/shop/products/{product.pk}/upload_video/"
            if backend_base
            else ""
        )

        payload: dict = {
            "backend": {
                "url": backend_url,
                "token": make_video_upload_token(product.pk),
            },
            "r2": None,
        }

        if is_r2_direct_upload_enabled():
            try:
                payload["r2"] = presign_video_upload(
                    product, content_type, filename
                )
            except Exception:
                logging.exception("prepare_video_upload presign failed")

        return Response(payload)

    @action(
        detail=True,
        methods=["post"],
        url_path="presign_video",
        permission_classes=[IsAdminUser],
    )
    def presign_video(
        self, request: Request, pk=None, slug=None, **kwargs
    ) -> Response:
        if not is_r2_direct_upload_enabled():
            return Response({"direct_upload": False, "reason": "r2_not_configured"})

        product = self.get_object()
        content_type = request.data.get("content_type") or "video/mp4"
        filename = request.data.get("filename") or ""

        try:
            upload = presign_video_upload(product, content_type, filename)
        except Exception:
            logging.exception("presign_video failed")
            return Response({"direct_upload": False})

        return Response({"direct_upload": True, "upload": upload})

    @action(
        detail=True,
        methods=["post"],
        url_path="confirm_video",
        permission_classes=[IsAdminUser],
    )
    def confirm_video(
        self, request: Request, pk=None, slug=None, **kwargs
    ) -> Response:
        if not is_r2_direct_upload_enabled():
            return Response(
                {"error": "Пряме завантаження недоступне"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        product = self.get_object()
        key = request.data.get("key")
        if not key or not validate_video_key(key, product):
            return Response(
                {"error": "Невірний ключ файлу"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if not object_exists(key):
            return Response(
                {"error": "Файл не знайдено в сховищі"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        title = request.data.get("title", "")
        order_num = int(request.data.get("order", product.videos.count()))
        pv = ProductVideo(product=product, title=title, order=order_num)
        pv.video.name = key
        pv.video._committed = True
        pv.save()

        return Response(
            ProductVideoSerializer(pv, context={"request": request}).data,
            status=status.HTTP_201_CREATED,
        )

    @action(
        detail=True,
        methods=["post"],
        url_path="upload_video",
        permission_classes=[IsAdminOrVideoUploadToken],
    )
    def upload_video(
        self, request: Request, pk=None, slug=None, **kwargs
    ) -> Response:
        product = self.get_object()
        if not (
            request.user.is_authenticated
            and request.user.is_staff
        ):
            token = request.headers.get("X-Video-Upload-Token", "")
            if not verify_video_upload_token(token, product.id):
                return Response(status=status.HTTP_403_FORBIDDEN)

        video_file = request.FILES.get("video")
        if not video_file:
            return Response(
                {"error": "Файл не завантажено"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        title = request.data.get("title", "")
        order_num = int(request.data.get("order", 0))
        pv = ProductVideo.objects.create(
            product=product, video=video_file, title=title, order=order_num
        )
        return Response(
            ProductVideoSerializer(pv, context={"request": request}).data,
            status=status.HTTP_201_CREATED,
        )

    @action(
        detail=True,
        methods=["delete"],
        url_path=r"delete_video/(?P<video_id>[^/.]+)",
        permission_classes=[IsAdminUser],
    )
    def delete_video(
        self, request: Request, pk=None, video_id=None, slug=None, **kwargs
    ) -> Response:
        product = self.get_object()
        try:
            ProductVideo.objects.get(pk=video_id, product=product).delete()
            return Response(status=status.HTTP_204_NO_CONTENT)
        except ProductVideo.DoesNotExist:
            return Response(status=status.HTTP_404_NOT_FOUND)

    @action(
        detail=True,
        methods=["post"],
        url_path="set_size",
        permission_classes=[IsAdminUser],
    )
    def set_size(
        self, request: Request, pk=None, slug=None, **kwargs
    ) -> Response:
        product = self.get_object()
        size = request.data.get("size")
        quantity = request.data.get("quantity", 0)
        if not size:
            return Response(
                {"error": "size обов'язковий"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        ps, created = ProductSize.objects.get_or_create(
            product=product,
            size=int(size),
            defaults={"quantity": int(quantity)},
        )
        if not created:
            ps.quantity = int(quantity)
            ps.save()
        return Response(
            ProductSizeListSerializer(ps).data,
            status=status.HTTP_201_CREATED if created else status.HTTP_200_OK,
        )

    @action(
        detail=True,
        methods=["post"],
        url_path="set_sizes",
        permission_classes=[IsAdminUser],
    )
    def set_sizes(
        self, request: Request, pk=None, slug=None, **kwargs
    ) -> Response:
        product = self.get_object()
        items = request.data.get("sizes")
        if not isinstance(items, list) or not items:
            return Response(
                {"error": "Передайте масив sizes з полями size та quantity"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        saved: list[ProductSize] = []
        for item in items:
            size = item.get("size")
            if size is None:
                continue
            quantity = int(item.get("quantity", 0))
            ps, created = ProductSize.objects.get_or_create(
                product=product,
                size=int(size),
                defaults={"quantity": quantity},
            )
            if not created:
                ps.quantity = quantity
                ps.save()
            saved.append(ps)

        return Response(
            ProductSizeListSerializer(saved, many=True).data,
            status=status.HTTP_200_OK,
        )

    @action(
        detail=True,
        methods=["delete"],
        url_path=r"delete_size/(?P<size_id>[^/.]+)",
        permission_classes=[IsAdminUser],
    )
    def delete_size(
        self, request: Request, pk=None, size_id=None, slug=None, **kwargs
    ) -> Response:
        product = self.get_object()
        try:
            ProductSize.objects.get(pk=size_id, product=product).delete()
            return Response(status=status.HTTP_204_NO_CONTENT)
        except ProductSize.DoesNotExist:
            return Response(status=status.HTTP_404_NOT_FOUND)


class VendorViewSet(viewsets.ModelViewSet):
    queryset = Vendor.objects.all()
    serializer_class = VendorSerializer
    permission_classes = [IsAdminOrReadOnly]


class WishlistViewSet(viewsets.ModelViewSet):
    serializer_class = WishlistSerializer
    permission_classes = [IsOwnerOrAdmin]
    pagination_class = None

    def get_queryset(self):
        return Wishlist.objects.filter(
            user=self.request.user
        ).prefetch_related("items__product__vendor", "items__product__sizes")

    @action(
        detail=False,
        methods=["post"],
        serializer_class=AddToWishlistSerializer,
        url_path="me/add_wish",
    )
    def add_item(self, request: Request) -> Response:
        wishlist, _ = Wishlist.objects.get_or_create(user=request.user)
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        product_id = serializer.validated_data["product"]

        try:
            product = Product.objects.get(id=product_id)
        except Product.DoesNotExist:
            return Response(
                {"error": "Product not found"},
                status=status.HTTP_404_NOT_FOUND,
            )

        WishlistItem.objects.get_or_create(wishlist=wishlist, product=product)
        return Response(
            {"message": "Product added to Wishlist!"},
            status=status.HTTP_201_CREATED,
        )

    @action(
        detail=False,
        methods=["post"],
        serializer_class=RemoveFromWishlistSerializer,
        url_path="me/remove_wish",
    )
    def remove_item(self, request: Request) -> Response:
        try:
            wishlist = Wishlist.objects.get(user=request.user)
        except Wishlist.DoesNotExist:
            return Response(
                {"error": "Wishlist not found"},
                status=status.HTTP_404_NOT_FOUND,
            )

        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        product_id = serializer.validated_data["product"]

        deleted, _ = WishlistItem.objects.filter(
            wishlist=wishlist, product_id=product_id
        ).delete()
        if deleted:
            return Response(status=status.HTTP_204_NO_CONTENT)
        return Response(
            {"error": "Product not found in Wishlist!"},
            status=status.HTTP_404_NOT_FOUND,
        )


class CartViewSet(viewsets.ModelViewSet):
    serializer_class = CartSerializer
    permission_classes = [IsOwnerOrAdmin]

    def get_queryset(self):
        return (
            Cart.objects.filter(user=self.request.user)
            .prefetch_related(
                "cart_items__product_size__product__vendor",
                "cart_items__product_size__product__images",
            )
            .order_by("-id")
        )

    @action(
        detail=False,
        methods=["post"],
        serializer_class=AddToCartSerializer,
        url_path="me/cart_add",
    )
    def add_item(self, request: Request) -> Response:
        cart, _ = Cart.objects.get_or_create(user=request.user)
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        product_size_id = serializer.validated_data["product_size"]
        quantity = serializer.validated_data.get("quantity", 1)
        product_size = ProductSize.objects.get(id=product_size_id)

        cart_item = CartItem.objects.filter(
            cart=cart, product_size=product_size
        ).first()
        current_quantity = cart_item.quantity if cart_item else 0
        if current_quantity + quantity > product_size.quantity:
            return Response(
                {"error": "Недостатньо товару в наявності для цього розміру."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        cart_item, created = CartItem.objects.get_or_create(
            cart=cart,
            product_size=product_size,
            defaults={"quantity": quantity},
        )
        if not created:
            cart_item.quantity += quantity
            cart_item.save()

        return Response(
            {"message": "Item added to cart"}, status=status.HTTP_201_CREATED
        )

    @action(
        detail=False,
        methods=["post"],
        serializer_class=RemoveFromCartSerializer,
        url_path="me/cart_remove",
    )
    def remove_item(self, request: Request) -> Response:
        try:
            cart = Cart.objects.get(user=request.user)
        except Cart.DoesNotExist:
            return Response(
                {"error": "Cart not found"}, status=status.HTTP_404_NOT_FOUND
            )

        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        product_size_id = serializer.validated_data["product_size"]
        quantity = serializer.validated_data.get("quantity", None)

        try:
            cart_item = CartItem.objects.get(
                cart=cart, product_size_id=product_size_id
            )
        except CartItem.DoesNotExist:
            return Response(
                {"error": "Item not found in cart"},
                status=status.HTTP_404_NOT_FOUND,
            )

        if quantity is None or quantity >= cart_item.quantity:
            cart_item.delete()
            return Response(
                {"message": "Item removed from cart"},
                status=status.HTTP_204_NO_CONTENT,
            )

        cart_item.quantity -= quantity
        cart_item.save()
        return Response(
            {"message": f"{quantity} of items has been removed from cart"},
            status=status.HTTP_200_OK,
        )


class OrderViewSet(viewsets.ModelViewSet):
    serializer_class = OrderSerializer
    permission_classes = [IsOwnerOrAdmin]

    def create(self, request, *args, **kwargs):
        return Response(status=status.HTTP_405_METHOD_NOT_ALLOWED)

    def get_queryset(self):
        base = Order.objects.select_related("delivery", "user", "created_by").prefetch_related(
            "items__product_size__product__vendor"
        )
        if can_manage_orders(self.request.user):
            return base
        return base.filter(user=self.request.user)

    @action(detail=False, methods=["post"], url_path="me/create_order")
    @transaction.atomic
    def create_order(self, request: Request) -> Response:
        delivery_data = (
            request.data.get("delivery")
            if isinstance(request.data.get("delivery"), dict)
            else request.data
        )
        delivery_serializer = DeliveryInfoSerializer(data=delivery_data)
        if not delivery_serializer.is_valid():
            return Response(
                delivery_serializer.errors, status=status.HTTP_400_BAD_REQUEST
            )

        try:
            cart = Cart.objects.get(user=request.user)
        except Cart.DoesNotExist:
            return Response(
                {"error": "Cart not found"}, status=status.HTTP_404_NOT_FOUND
            )

        if not cart.cart_items.exists():
            return Response(
                {"error": "Cart is empty"}, status=status.HTTP_400_BAD_REQUEST
            )

        items_data = [
            {
                "product_size": cart_item.product_size_id,
                "quantity": cart_item.quantity,
            }
            for cart_item in cart.cart_items.select_related("product_size").all()
        ]

        try:
            order = _create_order_with_items(
                order_user=request.user,
                delivery_data=delivery_serializer.validated_data,
                items_data=items_data,
                send_emails=True,
            )
        except ValidationError as exc:
            return Response(
                {"error": str(exc.message)},
                status=status.HTTP_400_BAD_REQUEST,
            )

        cart.cart_items.all().delete()
        order.refresh_from_db()
        return Response(
            OrderSerializer(order).data, status=status.HTTP_201_CREATED
        )

    @action(
        detail=False,
        methods=["post"],
        url_path="staff/create_order",
        permission_classes=[IsStaffOrSeller],
    )
    @transaction.atomic
    def staff_create_order(self, request: Request) -> Response:
        serializer = StaffCreateOrderSerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

        data = serializer.validated_data
        customer_email = (data.get("customer_email") or "").strip()
        order_user = None
        if customer_email:
            order_user = User.objects.filter(email__iexact=customer_email).first()

        try:
            order = _create_order_with_items(
                order_user=order_user,
                delivery_data=data["delivery"],
                items_data=data["items"],
                created_by=request.user,
                initial_status=data.get("status"),
                send_emails=bool(order_user),
            )
        except ValidationError as exc:
            return Response(
                {"error": str(exc.message)},
                status=status.HTTP_400_BAD_REQUEST,
            )

        order.refresh_from_db()
        return Response(
            OrderSerializer(order).data, status=status.HTTP_201_CREATED
        )

    @action(
        detail=False,
        methods=["post"],
        url_path="staff/quick_sale",
        permission_classes=[IsStaffOrSeller],
    )
    @transaction.atomic
    def quick_sale(self, request: Request) -> Response:
        """Оформлення з QR цінника: один розмір, самовивіз, очікує підтвердження оплати."""
        serializer = QuickSaleSerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

        data = serializer.validated_data
        name = (data.get("recipient_full_name") or "").strip() or "Покупець в магазині"
        phone = (data.get("recipient_phone") or "").strip() or settings.STORE_SALE_PHONE

        delivery_data = {
            "recipient_full_name": name,
            "recipient_phone": phone,
            "delivery_type": DeliveryInfo.DeliveryTypeChoices.PICKUP,
        }
        delivery_serializer = DeliveryInfoSerializer(data=delivery_data)
        delivery_serializer.is_valid(raise_exception=True)

        try:
            order = _create_order_with_items(
                order_user=None,
                delivery_data=delivery_serializer.validated_data,
                items_data=[
                    {
                        "product_size": data["product_size"],
                        "quantity": data["quantity"],
                    }
                ],
                created_by=request.user,
                initial_status=Order.StatusChoices.PENDING,
                send_emails=False,
                sale_channel=Order.SaleChannelChoices.STORE,
            )
        except ValidationError as exc:
            return Response(
                {"error": str(exc.message)},
                status=status.HTTP_400_BAD_REQUEST,
            )

        order.refresh_from_db()
        return Response(
            OrderSerializer(order).data, status=status.HTTP_201_CREATED
        )

    @action(
        detail=True,
        methods=["post"],
        url_path="confirm_payment",
        permission_classes=[IsStaffOrSeller],
    )
    def confirm_payment(self, request: Request, pk=None) -> Response:
        """Продавець підтверджує оплату (готівка або термінал) для замовлення з магазину."""
        order = self.get_object()
        if order.sale_channel != Order.SaleChannelChoices.STORE:
            return Response(
                {"error": "Підтвердження оплати лише для продажів у магазині"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if order.status != Order.StatusChoices.PENDING:
            return Response(
                {"error": "Замовлення вже оброблене"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        order.status = Order.StatusChoices.RECEIVED
        order.save()
        return Response(OrderSerializer(order).data, status=status.HTTP_200_OK)

    @action(detail=True, methods=["patch"], url_path="update_status")
    def update_status(self, request: Request, pk=None) -> Response:
        if not can_manage_orders(request.user):
            return Response(status=status.HTTP_403_FORBIDDEN)
        order = self.get_object()
        serializer = OrderStatusSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        new_status = serializer.validated_data["status"]
        order.status = new_status
        order.save()  # сигнали: processing→ТТН, cancelled→видалення ТТН + відновлення залишків

        response_data = OrderSerializer(order).data

        def _notify_status():
            try:
                send_order_status_update(order)
            except Exception:
                logger.exception("Order #%s: failed to send status update email", order.pk)

        transaction.on_commit(_notify_status)
        return Response(response_data, status=status.HTTP_200_OK)

    @action(
        detail=False,
        methods=["get"],
        url_path="pending_count",
        permission_classes=[IsStaffOrSeller],
    )
    def pending_count(self, request: Request) -> Response:
        count = Order.objects.filter(
            status=Order.StatusChoices.PENDING
        ).count()
        return Response({"count": count})

    @action(
        detail=True,
        methods=["post"],
        url_path="sync_np",
        permission_classes=[IsStaffOrSeller],
    )
    def sync_np_status(self, request: Request, pk=None) -> Response:
        """
        Запитує НП по ТТН і автоматично оновлює статус замовлення:
          • StatusCode 9  (Вручено)              → received
          • StatusCode 10 (Відмова)              → refused
          • StatusCode 11/101/102/14 (Повернення)→ refused
        """
        order = self.get_object()
        delivery = getattr(order, "delivery", None)

        if not delivery or not delivery.tracking_number:
            return Response(
                {"error": "Замовлення не має ТТН для відстеження"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        np_data = get_tracking_status(delivery.tracking_number)
        if not np_data:
            return Response(
                {"error": "Не вдалося отримати статус від Нової Пошти"},
                status=status.HTTP_502_BAD_GATEWAY,
            )

        np_code = str(np_data.get("StatusCode", ""))
        np_status_label = np_data.get("Status", "")

        RECEIVED_CODES = {"9"}
        REFUSED_CODES = {"10", "11", "14", "101", "102"}

        update_fields = []

        if np_code in RECEIVED_CODES:
            order.status = Order.StatusChoices.RECEIVED
            update_fields.append("status")
        elif np_code in REFUSED_CODES:
            order.status = Order.StatusChoices.REFUSED
            update_fields.append("status")

        if update_fields:
            order.save(update_fields=update_fields)

        return Response(
            {
                "np_code": np_code,
                "np_status": np_status_label,
                "order_status": order.status,
                "updated": bool(update_fields),
            }
        )

    @action(
        detail=True,
        methods=["post"],
        url_path="cancel",
        permission_classes=[IsAuthenticated],
    )
    def cancel_order(self, request: Request, pk=None) -> Response:
        order = self.get_object()
        if order.status != Order.StatusChoices.PENDING:
            return Response(
                {
                    "error": "Можна скасувати лише замовлення зі статусом 'Очікує обробки'."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )
        order.status = Order.StatusChoices.CANCELLED
        order.save()  # сигнали: _restore_stock_on_cancel, _delete_ttn_on_cancel

        response_data = OrderSerializer(order).data

        def _notify_cancel():
            try:
                send_order_status_update(order)
            except Exception:
                logger.exception("Order #%s: failed to send cancellation email", order.pk)

        transaction.on_commit(_notify_cancel)
        return Response(response_data, status=status.HTTP_200_OK)


class ProductSizeViewSet(viewsets.GenericViewSet):
    """QR цінники та контекст сканування для продажу в магазині."""

    queryset = ProductSize.objects.select_related("product__vendor").prefetch_related(
        "product__images"
    )
    permission_classes = [IsStaffOrSeller]

    @action(detail=True, methods=["get"], url_path="scan_info")
    def scan_info(self, request: Request, pk=None) -> Response:
        product_size = get_object_or_404(self.queryset, pk=pk)
        product = product_size.product
        if product.full_price <= 0:
            return Response(
                {"error": "Товар без ціни"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        main_image = product.images.filter(is_main=True).first() or product.images.first()
        image_url = None
        if main_image and main_image.image:
            image_url = request.build_absolute_uri(main_image.image.url)

        return Response(
            {
                "id": product_size.id,
                "size": product_size.size,
                "quantity": product_size.quantity,
                "vendor": product.vendor.name,
                "model_name": product.model_name,
                "prod_type": product.prod_type,
                "slug": product.slug,
                "discounted_price": str(product.discounted_price),
                "main_image": image_url,
                "scan_url": f"{settings.SITE_BASE_URL}/seller/scan/{product_size.id}",
            }
        )

    @action(detail=True, methods=["get"], url_path="qr_code")
    def qr_code(self, request: Request, pk=None) -> Response:
        import qrcode
        from PIL import Image, ImageDraw, ImageFont, ImageOps

        product_size = get_object_or_404(self.queryset, pk=pk)
        product = product_size.product
        scan_url = f"{settings.SITE_BASE_URL}/seller/scan/{product_size.id}"

        def load_font(size: int, bold: bool = False):
            candidates = (
                "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
                if bold
                else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
            )
            for path in candidates:
                try:
                    return ImageFont.truetype(path, size)
                except OSError:
                    continue
            return ImageFont.load_default()

        def centered_box_text(
            draw: ImageDraw.ImageDraw,
            box_x: int,
            y: int,
            box_width: int,
            text: str,
            font: ImageFont.ImageFont,
            fill: tuple[int, int, int] = (17, 24, 39),
        ) -> int:
            bbox = draw.textbbox((0, 0), text, font=font)
            width = bbox[2] - bbox[0]
            height = bbox[3] - bbox[1]
            draw.text((box_x + (box_width - width) / 2, y), text, font=font, fill=fill)
            return y + height

        def fit_font(
            text: str,
            start_size: int,
            max_width: int,
            bold: bool = False,
            min_size: int = 20,
        ) -> ImageFont.ImageFont:
            size = start_size
            while size > min_size:
                font = load_font(size, bold=bold)
                bbox = draw.textbbox((0, 0), text, font=font)
                if bbox[2] - bbox[0] <= max_width:
                    return font
                size -= 2
            return load_font(min_size, bold=bold)

        # 80x50 mm at 300 DPI.
        canvas_width = 945
        canvas_height = 591
        padding = 30
        canvas = Image.new("RGB", (canvas_width, canvas_height), "white")
        draw = ImageDraw.Draw(canvas)

        content_height = canvas_height - padding * 2
        image_box = (330, content_height)
        image_x = padding
        image_y = padding
        main_image = product.images.filter(is_main=True).first() or product.images.first()
        if main_image and main_image.image:
            try:
                with main_image.image.open("rb") as image_file:
                    product_img = Image.open(image_file).convert("RGB")
                    product_img = ImageOps.fit(
                        product_img,
                        image_box,
                        method=Image.Resampling.LANCZOS,
                        centering=(0.5, 0.5),
                    )
                    canvas.paste(product_img, (image_x, image_y))
            except Exception:
                logger.exception("Failed to render product image in QR label")

        qr_size = 370
        qr_x = canvas_width - padding - qr_size
        text_x = image_x + image_box[0] + 18
        text_width = qr_x - 18 - text_x
        price_value = str(int(product.discounted_price or product.full_price))
        text_lines = [
            (
                product.vendor.name,
                fit_font(product.vendor.name, 58, text_width, bold=True, min_size=30),
            ),
            (
                product.model_name,
                fit_font(product.model_name, 54, text_width, bold=True, min_size=30),
            ),
            ("SIZE", fit_font("SIZE", 42, text_width, bold=True, min_size=30)),
            (
                str(product_size.size),
                fit_font(str(product_size.size), 98, text_width, bold=True, min_size=54),
            ),
            (price_value, fit_font(price_value, 74, text_width, bold=True, min_size=46)),
            ("UAH", fit_font("UAH", 40, text_width, bold=True, min_size=30)),
        ]
        heights = [
            draw.textbbox((0, 0), text, font=font)[3]
            - draw.textbbox((0, 0), text, font=font)[1]
            for text, font in text_lines
        ]
        gap = max(2, (content_height - sum(heights)) / (len(text_lines) - 1))
        current_y = padding
        for index, (text, font) in enumerate(text_lines):
            current_y = centered_box_text(
                draw, text_x, int(current_y), text_width, text, font
            )
            if index < len(text_lines) - 1:
                current_y += gap

        qr = qrcode.QRCode(border=1, box_size=16)
        qr.add_data(scan_url)
        qr.make(fit=True)
        qr_img = qr.make_image(fill_color="black", back_color="white").convert("RGB")
        qr_img = ImageOps.contain(qr_img, (qr_size, qr_size), method=Image.Resampling.NEAREST)
        qr_y = (canvas_height - qr_img.height) // 2
        canvas.paste(qr_img, (qr_x, qr_y))

        buffer = io.BytesIO()
        canvas.save(buffer, format="PNG", dpi=(300, 300))
        return HttpResponse(buffer.getvalue(), content_type="image/png")


class ReviewViewSet(viewsets.ModelViewSet):
    serializer_class = ReviewSerializer

    def get_permissions(self):
        if self.action in ("list", "retrieve"):
            return [AllowAny()]
        if self.action == "create":
            return [IsAuthenticated()]
        return [IsOwnerOrAdmin()]

    def get_queryset(self):
        qs = Review.objects.select_related("user")
        product_id = self.request.query_params.get("product")
        if product_id:
            qs = qs.filter(product_id=product_id)
        return qs

    def perform_create(self, serializer):
        serializer.save(user=self.request.user)

    def create(self, request, *args, **kwargs):
        if Review.objects.filter(
            product_id=request.data.get("product"), user=request.user
        ).exists():
            return Response(
                {"error": "Ви вже залишили відгук для цього товару."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        return super().create(request, *args, **kwargs)

    @action(
        detail=False,
        methods=["get"],
        url_path="my",
        permission_classes=[IsAuthenticated],
    )
    def my_review(self, request: Request) -> Response:
        product_id = request.query_params.get("product")
        if not product_id:
            return Response(
                {"error": "Потрібен параметр product."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        review = Review.objects.filter(
            product_id=product_id, user=request.user
        ).first()
        if not review:
            return Response(None, status=status.HTTP_200_OK)
        return Response(
            ReviewSerializer(review, context={"request": request}).data
        )


@api_view(["GET"])
@permission_classes([AllowAny])
def nova_poshta_cities(request: Request) -> Response:
    query = request.query_params.get("q", "").strip()
    if len(query) < 2:
        return Response([], status=status.HTTP_200_OK)

    return Response(get_cities_list(query), status=status.HTTP_200_OK)


@api_view(["GET"])
@permission_classes([AllowAny])
def nova_poshta_warehouses(request: Request) -> Response:
    city_ref = request.query_params.get("city_ref", "").strip()
    query = request.query_params.get("q", "").strip()
    warehouse_type = request.query_params.get("type", "warehouse").strip()
    if not city_ref:
        return Response([], status=status.HTTP_200_OK)

    return Response(
        get_warehouses_list(city_ref, query, warehouse_type),
        status=status.HTTP_200_OK,
    )


@api_view(["GET"])
@permission_classes([IsAdminUser])
def manager_stats(request: Request) -> Response:
    today = timezone.now().date()

    def revenue_since(since_date):
        result = Order.objects.filter(
            created_at__date__gte=since_date
        ).aggregate(total=Sum("total_price"))["total"]
        return float(result or 0)

    top_products = list(
        OrderItem.objects.values(
            "product_size__product__id",
            "product_size__product__model_name",
            "product_size__product__vendor__name",
        )
        .annotate(
            total_sold=Sum("quantity"),
            total_revenue=Sum("price"),
        )
        .order_by("-total_sold")[:5]
    )

    orders_by_status = dict(
        Order.objects.values("status")
        .annotate(count=Count("id"))
        .values_list("status", "count")
    )

    return Response(
        {
            "revenue": {
                "today": revenue_since(today),
                "week": revenue_since(today - timedelta(days=7)),
                "month": revenue_since(today - timedelta(days=30)),
            },
            "orders_by_status": orders_by_status,
            "top_products": top_products,
            "total_orders": Order.objects.count(),
            "total_revenue": float(
                Order.objects.aggregate(total=Sum("total_price"))["total"] or 0
            ),
        }
    )


ACTIVE_STATUSES = [
    Order.StatusChoices.PENDING,
    Order.StatusChoices.PROCESSING,
    Order.StatusChoices.COMPLETED,
    Order.StatusChoices.RECEIVED,
]


def _get_report_queryset(period: str):
    today = timezone.now().date()
    if period == "week":
        from_date = today - timedelta(days=7)
    elif period == "month":
        from_date = today - timedelta(days=30)
    else:
        from_date = today

    return (
        OrderItem.objects.filter(
            order__created_at__date__gte=from_date,
            order__status__in=ACTIVE_STATUSES,
        )
        .values(
            vendor=F("product_size__product__vendor__name"),
            model_name=F("product_size__product__model_name"),
            size=F("product_size__size"),
            remaining=F("product_size__quantity"),
        )
        .annotate(
            sold_qty=Sum("quantity"),
            total_amount=Sum(
                ExpressionWrapper(
                    F("price") * F("quantity"),
                    output_field=DjDecimalField(
                        max_digits=12, decimal_places=2
                    ),
                )
            ),
        )
        .order_by("vendor", "model_name", "size")
    )


@api_view(["GET"])
@permission_classes([IsAdminUser])
def sales_report(request: Request) -> Response:
    period = request.GET.get("period", "today")
    rows = list(_get_report_queryset(period))
    return Response(
        [
            {
                "vendor": r["vendor"],
                "model_name": r["model_name"],
                "size": r["size"],
                "sold_qty": r["sold_qty"],
                "remaining": r["remaining"],
                "total_amount": float(r["total_amount"] or 0),
            }
            for r in rows
        ]
    )


@api_view(["GET"])
@permission_classes([IsAdminUser])
def sales_report_xlsx(request: Request) -> HttpResponse:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    period = request.GET.get("period", "today")
    rows = list(_get_report_queryset(period))

    period_labels = {
        "today": "сьогодні",
        "week": "за тиждень",
        "month": "за місяць",
    }
    period_label = period_labels.get(period, period)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Звіт"

    headers = [
        "Бренд",
        "Модель",
        "Розмір",
        "Продано, шт.",
        "Залишок, шт.",
        "Сума, грн.",
    ]
    header_fill = PatternFill(fill_type="solid", fgColor="4F46E5")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, r in enumerate(rows, 2):
        ws.cell(row=row_idx, column=1, value=r["vendor"])
        ws.cell(row=row_idx, column=2, value=r["model_name"])
        ws.cell(row=row_idx, column=3, value=r["size"])
        ws.cell(row=row_idx, column=4, value=r["sold_qty"])
        ws.cell(row=row_idx, column=5, value=r["remaining"])
        ws.cell(row=row_idx, column=6, value=float(r["total_amount"] or 0))
        if row_idx % 2 == 0:
            row_fill = PatternFill(fill_type="solid", fgColor="EEF2FF")
            for col in range(1, 7):
                ws.cell(row=row_idx, column=col).fill = row_fill

    # Total row
    total_row = len(rows) + 2
    ws.cell(row=total_row, column=1, value="РАЗОМ").font = Font(bold=True)
    ws.cell(
        row=total_row, column=4, value=sum(r["sold_qty"] or 0 for r in rows)
    ).font = Font(bold=True)
    ws.cell(
        row=total_row,
        column=6,
        value=sum(float(r["total_amount"] or 0) for r in rows),
    ).font = Font(bold=True)

    # Column widths
    for col, width in zip(range(1, 7), [20, 25, 10, 14, 14, 14]):
        ws.column_dimensions[
            ws.cell(row=1, column=col).column_letter
        ].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"report_{period}.xlsx"
    response = HttpResponse(
        buffer.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# ── Site visit tracking ────────────────────────────────────────────────────────


def _get_client_ip(request: Request) -> str:
    x_forwarded_for = request.META.get("HTTP_X_FORWARDED_FOR")
    if x_forwarded_for:
        return x_forwarded_for.split(",")[0].strip()
    return request.META.get("REMOTE_ADDR", "0.0.0.0")


@api_view(["POST"])
@permission_classes([AllowAny])
def track_visit(request: Request) -> Response:
    """Record a unique visit (one per IP per day)."""
    ip = _get_client_ip(request)
    today = timezone.now().date()
    SiteVisit.objects.get_or_create(ip=ip, date=today)
    return Response({"ok": True})


@api_view(["GET"])
@permission_classes([IsAdminUser])
def visit_stats(request: Request) -> Response:
    """Return visit and user counts for the manager dashboard."""
    from django.contrib.auth import get_user_model

    User = get_user_model()

    today = timezone.now().date()
    week_ago = today - timedelta(days=7)
    month_ago = today - timedelta(days=30)

    return Response(
        {
            "visits_today": SiteVisit.objects.filter(date=today).count(),
            "visits_week": SiteVisit.objects.filter(
                date__gte=week_ago
            ).count(),
            "visits_month": SiteVisit.objects.filter(
                date__gte=month_ago
            ).count(),
            "total_users": User.objects.filter(is_staff=False).count(),
        }
    )
